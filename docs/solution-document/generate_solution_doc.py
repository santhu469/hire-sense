"""
Generates docs/solution-document/HireSense-Solution-Document.docx

Source-of-truth script for the HireSense Solution Document. Edit this file
(not the .docx directly) and re-run it to regenerate the document, so the
document and its source stay in sync in version control.

Run from the repo root with the project venv:
    docs/.venv/bin/python docs/solution-document/generate_solution_doc.py
"""

import datetime
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parents[2]
DIAGRAMS = ROOT / "docs" / "architecture" / "diagrams"
OUT_FILE = Path(__file__).resolve().parent / "HireSense-Solution-Document.docx"

NAVY = RGBColor(0x1F, 0x4E, 0x79)
GRAY = RGBColor(0x59, 0x59, 0x59)

TODAY = datetime.date.today().strftime("%d %B %Y")


# --------------------------------------------------------------------------
# Low-level helpers
# --------------------------------------------------------------------------

def set_cell_background(cell, hex_color):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def add_page_break(doc):
    doc.add_page_break()


def add_title_page(doc):
    doc.add_paragraph().add_run().add_break()
    doc.add_paragraph().add_run().add_break()
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = t.add_run("HireSense")
    run.bold = True
    run.font.size = Pt(44)
    run.font.color.rgb = NAVY

    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = s.add_run("AI-Powered Candidate Evaluation Platform")
    run.font.size = Pt(20)
    run.font.color.rgb = GRAY

    s2 = doc.add_paragraph()
    s2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = s2.add_run("Solution Document")
    run.italic = True
    run.font.size = Pt(16)

    for _ in range(6):
        doc.add_paragraph()

    meta_rows = [
        ("Document status", "Draft v1.0"),
        ("Date", TODAY),
        ("Prepared for", "HireSense Client"),
        ("Scope", "Solution architecture & delivery plan (pre-implementation)"),
        ("Deployment target", "AWS (single-org, single-region, v1)"),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for label, value in meta_rows:
        row = table.add_row()
        row.cells[0].width = Inches(2.0)
        row.cells[1].width = Inches(3.5)
        row.cells[0].paragraphs[0].add_run(label).bold = True
        row.cells[1].paragraphs[0].add_run(value)
    add_page_break(doc)


def add_toc(doc, sections):
    doc.add_heading("Table of Contents", level=1)
    for num, title in sections:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.add_run(f"{num}.  {title}")
    add_page_break(doc)


def h1(doc, text):
    doc.add_heading(text, level=1)


def h2(doc, text):
    doc.add_heading(text, level=2)


def body(doc, text, italic=False, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = italic
    run.bold = bold
    return p


def bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def numbered(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Number")


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.autofit = True
    hdr_cells = table.rows[0].cells
    for i, htext in enumerate(headers):
        hdr_cells[i].text = ""
        run = hdr_cells[i].paragraphs[0].add_run(htext)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_background(hdr_cells[i], "1F4E79")
    for row_data in rows:
        row_cells = table.add_row().cells
        for i, val in enumerate(row_data):
            row_cells[i].text = str(val)
    doc.add_paragraph()
    return table


def add_figure(doc, filename, caption, width_in=None, height_in=None):
    path = DIAGRAMS / filename
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    if height_in:
        run.add_picture(str(path), height=Inches(height_in))
    else:
        run.add_picture(str(path), width=Inches(width_in))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_run = cap.add_run(caption)
    cap_run.italic = True
    cap_run.font.size = Pt(10)
    cap_run.font.color.rgb = GRAY
    doc.add_paragraph()


# --------------------------------------------------------------------------
# Document assembly
# --------------------------------------------------------------------------

def build():
    doc = Document()

    # base font
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    add_title_page(doc)

    sections = [
        ("1", "Purpose & Business Context"),
        ("2", "Scope"),
        ("3", "Stakeholders & Roles"),
        ("4", "Functional Requirements Traceability"),
        ("5", "Non-Functional Requirements"),
        ("6", "Solution Architecture Overview"),
        ("7", "Component Design"),
        ("8", "Data Model"),
        ("9", "Primary Application / Process Flow"),
        ("10", "AWS Deployment Architecture"),
        ("11", "Technology Stack Summary"),
        ("12", "Security, Compliance & PII Handling"),
        ("13", "Phased Delivery Roadmap"),
        ("14", "Open Questions & Assumptions"),
        ("15", "Appendix: Glossary"),
    ]
    add_toc(doc, sections)

    # 1. Purpose & Business Context ----------------------------------------
    h1(doc, "1. Purpose & Business Context")
    body(
        doc,
        "HireSense is an AI-powered candidate evaluation application designed to reduce the "
        "manual effort of resume screening. It allows recruiters and hiring managers to create "
        "or upload Job Descriptions (JDs), bulk-upload candidate resumes against a JD, and have "
        "every candidate automatically scored and explained against that JD. The system is built "
        "to handle 10-100+ resumes per JD.",
    )
    body(doc, "The core operating principle that governs every design decision in this document:")
    q = doc.add_paragraph()
    q.paragraph_format.left_indent = Inches(0.4)
    qr = q.add_run(
        "“Use AI to reduce resume-screening effort and provide structured, explainable "
        "recommendations while keeping the final hiring decision with the human recruiter or "
        "hiring manager.”"
    )
    qr.italic = True
    qr.font.color.rgb = NAVY
    body(
        doc,
        "In practice this means HireSense is a decision-support system, not a decision-making "
        "system: the AI never finalizes a candidate's status, and no candidate-facing email is "
        "ever sent without an explicit human approval action.",
    )

    # 2. Scope ---------------------------------------------------------------
    h1(doc, "2. Scope")
    h2(doc, "2.1 In scope — Phase 1 (this document)")
    bullets(
        doc,
        [
            "JD creation (manual, upload, AI-generate/improve) and management of multiple active JDs.",
            "Bulk candidate resume upload and parsing, associated to a JD.",
            "Configurable, weighted AI evaluation producing an overall score, category scores, and "
            "explainable insights per candidate.",
            "Candidate ranking, Top-10 identification, filtering and sorting.",
            "Human-in-the-loop decisioning (Shortlisted / Hold / Rejected) with full decision history.",
            "Re-evaluation / rescoring with version history when JD, criteria, weights, or candidate "
            "data changes.",
            "AI-drafted candidate email communication, sent only after explicit human approval.",
            "Role-based access control: Admin, Recruiter/Evaluator, Viewer/Hiring Manager.",
            "Single organization (no multi-tenancy) deployment on AWS.",
        ],
    )
    h2(doc, "2.2 Explicitly out of scope for Phase 1")
    bullets(
        doc,
        [
            "Multi-tenant / multi-organization support (single-org for v1; data model leaves room "
            "to add an Organization entity later without a rewrite).",
            "SSO / enterprise identity provider integration (email/password auth for v1).",
            "Applicant-facing self-service portal (candidates do not log in; they are only recipients "
            "of outbound email).",
            "Interview scheduling, offer management, or downstream HRIS/ATS integrations.",
        ],
    )

    # 3. Stakeholders & Roles --------------------------------------------
    h1(doc, "3. Stakeholders & Roles")
    add_table(
        doc,
        ["Role", "Can do"],
        [
            (
                "Admin",
                "Manage users and entitlements; configure application-level settings "
                "(default evaluation weights, templates).",
            ),
            (
                "Recruiter / Evaluator",
                "Create JDs; upload candidates; trigger evaluation and rescoring; update candidate "
                "status; initiate and send candidate communication.",
            ),
            (
                "Viewer / Hiring Manager",
                "Read-only: review JDs, candidate evaluations, recommendations, and status. Cannot "
                "change status, weights, or send email.",
            ),
        ],
    )

    # 4. Functional Requirements Traceability -----------------------------
    h1(doc, "4. Functional Requirements Traceability")
    body(doc, "Every functional requirement in the Business Requirements Document maps to a named service in this architecture (see Section 7):")
    add_table(
        doc,
        ["BRD Section", "Requirement", "Owning Component"],
        [
            ("1", "Create/manage JDs using AI", "JD Service"),
            ("2 (Step 1)", "Create/upload/AI-generate JD; multiple active JDs", "JD Service"),
            ("2 (Step 2)", "Bulk resume upload; extract info; associate to JD; handle 100+/JD", "Resume Ingestion Service"),
            ("2 (Step 3)", "Overall + category scores; configurable weightage", "Evaluation Orchestrator"),
            ("2 (Step 4)", "Explainable insights (strengths, gaps, recommendation, ranking rationale)", "Evaluation Orchestrator"),
            ("3", "Rank candidates; Top 10; filter/sort", "Ranking & Shortlisting Service"),
            ("4", "Human sets Shortlisted/Hold/Rejected; decision history", "Decision & Audit Service"),
            ("5", "Rescoring on JD/criteria/weight/data change; version history", "Evaluation Orchestrator"),
            ("6", "AI-drafted email; human-approved send only", "Communication Service"),
            ("7", "RBAC: Admin / Recruiter / Viewer", "API Gateway / Auth layer, all services"),
        ],
    )

    # 5. Non-Functional Requirements ---------------------------------------
    h1(doc, "5. Non-Functional Requirements")
    add_table(
        doc,
        ["Category", "Requirement", "Design response"],
        [
            ("Scale", "10-100+ resumes per JD", "Bulk upload + async job queue (SQS) + horizontally scalable worker service; nothing resume-related runs synchronously in the request path."),
            ("Explainability", "Every score must come with a reason", "LLM evaluation calls always return structured insights (strengths, gaps, recommendation) alongside the numeric score — never a bare number."),
            ("Governance", "AI never finalizes a decision", "Status transitions are a separate, human-only write path (Decision & Audit Service); no code path lets an evaluation job set candidate status."),
            ("Auditability", "Decision and score history must be retained", "CandidateDecision is an append-only log; Evaluation rows are versioned (is_latest flag) rather than overwritten."),
            ("Control", "Email must require explicit human approval", "EmailDraft has a draft/sent state machine; the send action is a distinct, logged, human-triggered API call."),
            ("Configurability", "Evaluation weights must be adjustable", "EvaluationCriteria is a versioned, editable table per JD, not a hardcoded prompt constant."),
            ("Security", "Resumes contain PII", "Resumes stored in a private S3 bucket (SSE encryption), access scoped by JD/role, secrets never in code."),
        ],
    )

    # 6. Solution Architecture Overview -------------------------------------
    h1(doc, "6. Solution Architecture Overview")
    body(
        doc,
        "The system is a modular monolith at the application layer (FastAPI, Python) fronted by a "
        "Next.js web client, backed by PostgreSQL for structured data, S3 for resume files, and an "
        "SQS-backed job queue for asynchronous AI evaluation work. Services are organized as clean "
        "internal modules for v1; the boundaries are drawn so any module can be split into an "
        "independent deployable later without a data-model rewrite.",
    )
    add_figure(doc, "system_architecture.png", "Figure 1 — HireSense logical system architecture.", width_in=6.5)

    # 7. Component Design ----------------------------------------------------
    h1(doc, "7. Component Design")

    h2(doc, "7.1 JD Service")
    bullets(
        doc,
        [
            "CRUD for Job Descriptions; each JD is an independent evaluation workspace.",
            "AI actions: generate a JD from a title/bullet list, or improve/rewrite an uploaded JD "
            "— returns structured requirements (must-have skills, nice-to-have, experience range, "
            "responsibilities), not just free text, so it can be diffed and edited.",
        ],
    )

    h2(doc, "7.2 Resume Ingestion Service")
    bullets(
        doc,
        [
            "Bulk upload endpoint accepts PDF/DOCX; raw files land in S3; a parse job is enqueued "
            "per file — never parsed synchronously in the request.",
            "Text extraction + LLM-based structured extraction (name, contact, skills, experience "
            "timeline, education, certifications) stored as a candidate profile.",
        ],
    )

    h2(doc, "7.3 Evaluation Orchestrator (core AI engine)")
    bullets(
        doc,
        [
            "Runs as background jobs (SQS-triggered) so 100+ resumes never block the API.",
            "Per-candidate LLM call uses structured output (JSON schema / tool-use), not free text, "
            "returning category scores and an insights object (summary, strengths, matching skills, "
            "gaps, differentiators, concerns, recommendation).",
            "The final weighted overall score is computed deterministically in application code from "
            "category scores × configured weights — the LLM never computes the weighted total, "
            "which keeps re-weighting cheap (no LLM re-call needed for a weight-only change).",
            "‘Why ranked above/below’ is generated by comparing category scores/insights between "
            "neighboring-ranked candidates.",
        ],
    )

    h2(doc, "7.4 Ranking & Shortlisting Service")
    bullets(
        doc,
        [
            "Pure query logic over Evaluation rows — sort/filter by score, category, experience, "
            "skills, status. No AI call on the read path.",
            "‘Top 10’ is simply the top-N by overall_score, surfaced prominently in the UI.",
        ],
    )

    h2(doc, "7.5 Decision & Audit Service")
    bullets(
        doc,
        [
            "Owns the Shortlisted / Hold / Rejected transitions — always human-initiated.",
            "CandidateDecision is append-only; current status is the latest row. This table is the "
            "system's audit trail for hiring decisions.",
        ],
    )

    h2(doc, "7.6 Re-Evaluation / Rescoring")
    bullets(
        doc,
        [
            "Triggered by JD edits, weight changes, criteria changes, or new candidate information.",
            "Evaluation rows are never overwritten — a new version is inserted and is_latest is "
            "flipped, so history is preserved and the UI always shows the latest score.",
            "Weight-only changes recompute the weighted sum without a new LLM call; JD/criteria "
            "definition changes trigger a full re-evaluation.",
        ],
    )

    h2(doc, "7.7 Communication Service")
    bullets(
        doc,
        [
            "LLM drafts email content (subject/body) for shortlist, hold, rejection, or "
            "interview-next-step — saved as an EmailDraft in draft state.",
            "No code path calls the email provider without a human-approval record attached to the "
            "request; the UI requires an explicit Send action.",
        ],
    )

    h2(doc, "7.8 RBAC / Auth")
    bullets(
        doc,
        [
            "Email/password authentication; JWT access/refresh tokens.",
            "Role stored on the User record; permission checks enforced at the API route/action level "
            "(e.g., only Recruiter/Admin can change status or send email; Viewer is read-only).",
        ],
    )

    # 8. Data Model -----------------------------------------------------------
    h1(doc, "8. Data Model")
    add_figure(doc, "data_model_er.png", "Figure 2 — HireSense core entity-relationship model.", width_in=6.5)
    body(
        doc,
        "Notes: Evaluation and CandidateDecision are intentionally append-only/versioned tables "
        "(not update-in-place) to satisfy the BRD's history requirements for rescoring and decisions. "
        "JSONB columns (structured_requirements, parsed_profile, category_scores, insights) hold "
        "semi-structured AI output so the schema doesn't need to change every time a new evaluation "
        "category is added — only EvaluationCriteria (a normal relational table) needs a new row.",
    )

    # 9. Process Flow -----------------------------------------------------
    h1(doc, "9. Primary Application / Process Flow")
    add_figure(doc, "process_flow.png", "Figure 3 — Primary application flow with the two human-in-the-loop gates.", width_in=6.5)
    body(
        doc,
        "The two boxed gates in Figure 3 are hard product requirements, not defaults that can be "
        "configured away: a candidate's status can only be set by an authenticated human action, and "
        "an email can only leave the system after an explicit human send click on a specific draft.",
    )

    # 10. AWS Deployment -----------------------------------------------------
    h1(doc, "10. AWS Deployment Architecture")
    body(
        doc,
        "Single-region, single-org deployment. Frontend and backend are independently deployable "
        "containers on ECS Fargate behind one Application Load Balancer (path-based routing); "
        "evaluation jobs run in a separate Fargate worker service so a burst of 100+ resume uploads "
        "cannot starve interactive API traffic.",
    )
    add_figure(doc, "aws_deployment.png", "Figure 4 — HireSense AWS deployment topology.", height_in=8.5)
    h2(doc, "10.1 Service-by-service rationale")
    add_table(
        doc,
        ["AWS Service", "Why"],
        [
            ("Route 53 + ACM + ALB", "DNS + managed TLS termination + path-based routing to frontend vs. backend target groups."),
            ("ECS Fargate (frontend)", "Runs `next start` (SSR) as a container; no server management, scales on request load."),
            ("ECS Fargate (backend API)", "FastAPI app; stateless, horizontally scalable behind the ALB."),
            ("ECS Fargate (evaluation worker)", "Consumes the SQS queue; scales independently of API traffic based on queue depth."),
            ("RDS PostgreSQL (Multi-AZ)", "Primary datastore for JDs, candidates, evaluations, decisions, users — Multi-AZ for availability."),
            ("S3", "Private bucket for resume files and (optionally) static frontend assets; server-side encryption."),
            ("SQS", "Decouples resume upload/JD-change events from the evaluation workers; absorbs bulk-upload bursts."),
            ("Secrets Manager", "OpenAI API key, DB credentials, JWT signing key — never stored in code or images."),
            ("CloudWatch", "Centralized logs/metrics for all three Fargate services; basis for cost and latency dashboards on LLM calls."),
            ("ECR", "Private container registry; source for all three Fargate task definitions."),
            ("SES", "Outbound candidate email delivery, only invoked after human approval."),
        ],
    )

    # 11. Tech stack ----------------------------------------------------------
    h1(doc, "11. Technology Stack Summary")
    add_table(
        doc,
        ["Layer", "Choice"],
        [
            ("Frontend", "Next.js (TypeScript), Tailwind CSS, TanStack Table for the ranking grid"),
            ("Backend", "Python, FastAPI, SQLAlchemy + Alembic migrations"),
            ("Auth", "Email/password, bcrypt password hashing, JWT access/refresh tokens"),
            ("Database", "PostgreSQL (AWS RDS, Multi-AZ)"),
            ("Object storage", "AWS S3 (resume files)"),
            ("Async jobs", "AWS SQS + a dedicated Python worker service (boto3)"),
            ("AI / LLM", "OpenAI API, structured outputs / function calling for scoring and extraction"),
            ("Email delivery", "AWS SES"),
            ("Compute / hosting", "AWS ECS Fargate (frontend, backend API, worker — 3 services)"),
            ("Secrets", "AWS Secrets Manager"),
            ("Observability", "AWS CloudWatch (logs, metrics)"),
            ("CI/CD", "GitHub Actions → build & push to ECR → deploy to ECS (future work)"),
        ],
    )

    # 12. Security ----------------------------------------------------------
    h1(doc, "12. Security, Compliance & PII Handling")
    bullets(
        doc,
        [
            "Resumes and parsed candidate profiles are PII — stored in a private S3 bucket and a "
            "PostgreSQL database that are not publicly reachable; access is scoped by role and JD.",
            "All data in transit uses TLS (ACM-terminated at the ALB); S3 and RDS use server-side "
            "encryption at rest.",
            "Secrets (LLM API key, DB credentials, JWT signing key) live in Secrets Manager, never in "
            "source control or container images.",
            "Every status change and every sent email is attributable to a specific authenticated "
            "user and timestamped — satisfies basic audit/compliance evidence needs.",
            "Data retention period for resumes/PII is an open question for the client (Section 14) — "
            "recommend a configurable retention/purge policy once confirmed.",
        ],
    )

    # 13. Roadmap -------------------------------------------------------------
    h1(doc, "13. Phased Delivery Roadmap")
    add_table(
        doc,
        ["Phase", "Scope", "Outcome"],
        [
            ("0. Discovery", "Finalize data model, confirm open questions, environment setup", "Signed-off architecture (this document)"),
            ("1. Core loop (MVP)", "Auth, manual JD create, single resume upload, AI evaluation (fixed weights), ranked list, manual status change", "Provable end-to-end value"),
            ("2. Scale ingestion", "Bulk upload, async job queue, progress UI, 100+ resume handling", "Meets stated scale requirement"),
            ("3. Configurability & history", "Configurable weights, versioned rescoring, decision audit trail", "Meets governance requirements"),
            ("4. RBAC", "Full 3-role model, permission enforcement", "Meets access control requirement"),
            ("5. Communication", "AI email drafting + human-approved send, all four templates", "Closes the loop end-to-end"),
            ("6. Polish / AI JD tools", "AI JD generation/improvement, candidate comparison rationale", "Full BRD coverage"),
        ],
    )

    # 14. Open questions -------------------------------------------------------
    h1(doc, "14. Open Questions & Assumptions")
    h2(doc, "14.1 Confirmed with client")
    bullets(
        doc,
        [
            "Single organization / single-tenant deployment (no multi-tenancy in v1).",
            "Authentication: email/password (no SSO in v1).",
            "Frontend: Next.js. Backend: Python (FastAPI). Deployment: AWS.",
        ],
    )
    h2(doc, "14.2 Still open — to confirm before Phase 1 build starts")
    bullets(
        doc,
        [
            "Data retention period for resumes/PII — how long is data kept, is there an obligation "
            "to purge (e.g., GDPR/CCPA-style)?",
            "Expected turnaround SLA for evaluating a 100-resume batch (seconds vs. minutes acceptable)?",
            "Email sending identity — does the client have an existing domain to verify in SES, or "
            "should HireSense provide a shared sending domain initially?",
            "Any compliance constraint on which region/vendor may process candidate PII through an "
            "LLM API?",
        ],
    )

    # 15. Glossary --------------------------------------------------------------
    h1(doc, "15. Appendix: Glossary")
    add_table(
        doc,
        ["Term", "Meaning"],
        [
            ("JD", "Job Description — an independent candidate-evaluation workspace in HireSense."),
            ("RBAC", "Role-Based Access Control."),
            ("Overall score", "The single 0-100 score for a candidate against a JD, computed as a weighted sum of category scores."),
            ("Category score", "A sub-score for one evaluation dimension, e.g. Skills Match, Domain Experience."),
            ("Human-in-the-loop", "A workflow step that requires an explicit action by an authenticated human user before it takes effect."),
            ("is_latest", "Flag on the Evaluation table marking the current version of a candidate's score after rescoring."),
        ],
    )

    OUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT_FILE))
    print(f"Wrote {OUT_FILE}")


if __name__ == "__main__":
    build()
